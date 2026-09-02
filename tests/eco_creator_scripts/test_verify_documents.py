import httpx
from docx import Document
from docx.shared import Emu

import verify_documents
from verify_documents import arena

TODAY = verify_documents.datetime.now(verify_documents.timezone.utc).date().isoformat()


def _build_docx(
    path,
    department="Quality",
    author="Preparer Person",
    department_head="Head Person",
    signature_fields=None,
    revisions=None,
    header_title="Detector Head Assembly",
    header_number="FO-00127",
    header_revision="02-A",
    footer_id_line="FO-00127 Rev 02-A",
    footer_eco_number="ECO-00042",
    footer_revision_date=None,
):
    """Mirrors dilon-document-compiler's generate_dilon_doc.py /
    dilon_docx_common.py table layout (populate_header/populate_footer,
    create_signature_table/create_revision_table) closely enough for the
    header-based table finder to locate all of them."""
    if footer_revision_date is None:
        footer_revision_date = TODAY
    signature_fields = signature_fields if signature_fields is not None else [
        {"department": "Quality", "name": "Jane Doe"}
    ]
    revisions = revisions if revisions is not None else [
        {"number": "01", "description": "Initial release", "eco_number": "ECO-00042", "eco_date": "2026-09-01"}
    ]

    doc = Document()
    section = doc.sections[0]

    # --- Header ---
    header_table = section.header.add_table(rows=1, cols=4, width=Emu(0))
    hrow = header_table.rows[0]
    hrow.cells[1].paragraphs[0].add_run(f"Title: {header_title}")
    p = hrow.cells[1].add_paragraph()
    p.add_run(f"Number: {header_number}")
    hrow.cells[2].paragraphs[0].add_run(f"Rev {header_revision}")

    # --- Footer ---
    footer_table = section.footer.add_table(rows=2, cols=3, width=Emu(0))
    frow = footer_table.rows[0]
    frow.cells[0].paragraphs[0].add_run(footer_id_line)
    frow.cells[1].paragraphs[0].add_run(footer_eco_number)
    frow.cells[2].paragraphs[0].add_run(f"Revision Date: {footer_revision_date}")

    # --- Signature table (body) ---
    sig_table = doc.add_table(rows=4 + len(signature_fields), cols=3)
    sig_table.rows[0].cells[0].text = "Group"
    sig_table.rows[0].cells[1].text = "Preparer"
    sig_table.rows[0].cells[2].text = "Signature"
    sig_table.rows[1].cells[0].text = department
    sig_table.rows[1].cells[1].text = author
    sig_table.rows[1].cells[2].text = "Electronic"
    sig_table.rows[2].cells[0].text = "Department"
    sig_table.rows[2].cells[1].text = "Name"
    sig_table.rows[2].cells[2].text = "Signature"
    sig_table.rows[3].cells[0].text = department
    sig_table.rows[3].cells[1].text = department_head
    sig_table.rows[3].cells[2].text = "Electronic"
    for i, entry in enumerate(signature_fields):
        row = sig_table.rows[4 + i]
        row.cells[0].text = entry["department"]
        row.cells[1].text = entry["name"]
        row.cells[2].text = "Electronic"

    doc.add_paragraph()  # separator, matches generator's real layout

    # --- Revision-history table (body) ---
    rev_table = doc.add_table(rows=2 + len(revisions), cols=4)
    title_cell = rev_table.rows[0].cells[0]
    for i in range(1, 4):
        title_cell.merge(rev_table.rows[0].cells[i])
    title_cell.text = "REVISION HISTORY"
    headers = ["REV #", "DESCRIPTION OF CHANGE", "ECO #", "DATE"]
    for i, h in enumerate(headers):
        rev_table.rows[1].cells[i].text = h
    for i, rev in enumerate(revisions):
        row = rev_table.rows[2 + i]
        row.cells[0].text = rev["number"]
        row.cells[1].text = rev["description"]
        row.cells[2].text = rev["eco_number"]
        row.cells[3].text = rev["eco_date"]

    doc.save(str(path))
    return path


def _mock_editions(arena_api, docx_dt="2026-09-01T12:00:00Z", pdf_dt="2026-09-01T12:05:00Z"):
    arena_api.get(f"{arena.ARENA_API_BASE}/files/F1/editions").mock(
        return_value=httpx.Response(200, json={"results": [{"creationDateTime": docx_dt}]})
    )
    arena_api.get(f"{arena.ARENA_API_BASE}/files/F2/editions").mock(
        return_value=httpx.Response(200, json={"results": [{"creationDateTime": pdf_dt}]})
    )


def _mock_history(arena_api, reviewer_names):
    arena_api.get(f"{arena.ARENA_API_BASE}/changes/CHG1/history").mock(
        return_value=httpx.Response(
            200,
            json={
                "results": [
                    {"action": "Add Additional Reviewer", "newValue": f"{n} (Approval Required)"}
                    for n in reviewer_names
                ]
            },
        )
    )


def _mock_change_and_item(
    arena_api,
    eco_number="ECO-00042",
    item_number="FO-00127",
    item_name="Detector Head Assembly",
):
    arena_api.get(f"{arena.ARENA_API_BASE}/changes/CHG1").mock(
        return_value=httpx.Response(200, json={"guid": "CHG1", "number": eco_number})
    )
    arena_api.get(f"{arena.ARENA_API_BASE}/items/ITEM1").mock(
        return_value=httpx.Response(200, json={"guid": "ITEM1", "number": item_number, "name": item_name})
    )


def _verify_one(docx_path, expected_revision="02-A", is_form=False):
    return verify_documents.verify(
        "CHG1",
        [
            {
                "item_guid": "ITEM1",
                "docx_path": str(docx_path),
                "docx_file_guid": "F1",
                "pdf_file_guid": "F2",
                "expected_revision": expected_revision,
                "is_form": is_form,
            }
        ],
    )


# ---- Section 1: header/footer format compliance ----


def test_matching_header_footer_is_clean(tmp_path, arena_api):
    docx_path = _build_docx(tmp_path / "doc.docx", department_head="Head Person")
    _mock_history(arena_api, ["Jane Doe", "Head Person"])
    _mock_editions(arena_api)
    _mock_change_and_item(arena_api)

    result = _verify_one(docx_path)

    assert result["documents"][0]["errors"] == []
    assert result["ok"] is True


def test_header_title_mismatch_is_blocking(tmp_path, arena_api):
    docx_path = _build_docx(tmp_path / "doc.docx", department_head="Head Person", header_title="Wrong Title")
    _mock_history(arena_api, ["Jane Doe", "Head Person"])
    _mock_editions(arena_api)
    _mock_change_and_item(arena_api)

    result = _verify_one(docx_path)

    assert result["ok"] is False
    assert any("Header Title" in e for e in result["documents"][0]["errors"])


def test_header_number_mismatch_is_blocking(tmp_path, arena_api):
    docx_path = _build_docx(tmp_path / "doc.docx", department_head="Head Person", header_number="FO-99999")
    _mock_history(arena_api, ["Jane Doe", "Head Person"])
    _mock_editions(arena_api)
    _mock_change_and_item(arena_api)

    result = _verify_one(docx_path)

    assert result["ok"] is False
    errors = result["documents"][0]["errors"]
    assert any("Header Number" in e for e in errors)
    assert any("does not appear anywhere in the header" in e for e in errors)


def test_header_revision_mismatch_is_blocking(tmp_path, arena_api):
    docx_path = _build_docx(tmp_path / "doc.docx", department_head="Head Person", header_revision="01")
    _mock_history(arena_api, ["Jane Doe", "Head Person"])
    _mock_editions(arena_api)
    _mock_change_and_item(arena_api)

    result = _verify_one(docx_path, expected_revision="02-A")

    assert result["ok"] is False
    assert any("Header Rev" in e for e in result["documents"][0]["errors"])


def test_footer_id_line_mismatch_is_blocking(tmp_path, arena_api):
    docx_path = _build_docx(
        tmp_path / "doc.docx", department_head="Head Person", footer_id_line="FO-00127 Rev 01"
    )
    _mock_history(arena_api, ["Jane Doe", "Head Person"])
    _mock_editions(arena_api)
    _mock_change_and_item(arena_api)

    result = _verify_one(docx_path)

    assert result["ok"] is False
    assert any("Footer ID line" in e for e in result["documents"][0]["errors"])


def test_footer_eco_number_mismatch_is_blocking(tmp_path, arena_api):
    docx_path = _build_docx(
        tmp_path / "doc.docx", department_head="Head Person", footer_eco_number="ECO-WRONG"
    )
    _mock_history(arena_api, ["Jane Doe", "Head Person"])
    _mock_editions(arena_api)
    _mock_change_and_item(arena_api)

    result = _verify_one(docx_path)

    assert result["ok"] is False
    assert any("Footer ECO #" in e for e in result["documents"][0]["errors"])


def test_footer_revision_date_not_today_is_blocking(tmp_path, arena_api):
    docx_path = _build_docx(
        tmp_path / "doc.docx", department_head="Head Person", footer_revision_date="2020-01-01"
    )
    _mock_history(arena_api, ["Jane Doe", "Head Person"])
    _mock_editions(arena_api)
    _mock_change_and_item(arena_api)

    result = _verify_one(docx_path)

    assert result["ok"] is False
    assert any("Revision Date" in e for e in result["documents"][0]["errors"])


def test_missing_header_table_is_an_error(tmp_path, arena_api):
    docx_path = _build_docx(tmp_path / "doc.docx", department_head="Head Person")
    # Wipe the header table out after building, to simulate a document
    # with no header table at all.
    doc = Document(str(docx_path))
    doc.sections[0].header.tables[0]._element.getparent().remove(
        doc.sections[0].header.tables[0]._element
    )
    doc.save(str(docx_path))

    _mock_history(arena_api, ["Jane Doe", "Head Person"])
    _mock_editions(arena_api)
    _mock_change_and_item(arena_api)

    result = _verify_one(docx_path)

    assert result["ok"] is False
    assert any("No header table found" in e for e in result["documents"][0]["errors"])


def _build_form_docx(path, **header_footer_kwargs):
    """Mirrors generate_dilon_form.py: running header/footer only, no
    signature or revision-history body tables at all."""
    doc = Document()
    section = doc.sections[0]

    header_table = section.header.add_table(rows=1, cols=4, width=Emu(0))
    hrow = header_table.rows[0]
    hrow.cells[1].paragraphs[0].add_run(f"Title: {header_footer_kwargs.get('header_title', 'Detector Head Assembly')}")
    p = hrow.cells[1].add_paragraph()
    p.add_run(f"Number: {header_footer_kwargs.get('header_number', 'FO-00127')}")
    hrow.cells[2].paragraphs[0].add_run(f"Rev {header_footer_kwargs.get('header_revision', '02-A')}")

    footer_table = section.footer.add_table(rows=2, cols=3, width=Emu(0))
    frow = footer_table.rows[0]
    frow.cells[0].paragraphs[0].add_run(header_footer_kwargs.get("footer_id_line", "FO-00127 Rev 02-A"))
    frow.cells[1].paragraphs[0].add_run(header_footer_kwargs.get("footer_eco_number", "ECO-00042"))
    frow.cells[2].paragraphs[0].add_run(f"Revision Date: {header_footer_kwargs.get('footer_revision_date', TODAY)}")

    doc.save(str(path))
    return path


# ---- 2a/2b applicability: skipped entirely for form documents ----


def test_form_document_skips_signature_and_revision_checks(tmp_path, arena_api):
    docx_path = _build_form_docx(tmp_path / "form.docx")
    _mock_history(arena_api, [])  # no reviewers recorded -- must not matter for a form
    _mock_change_and_item(arena_api)

    result = _verify_one(docx_path, is_form=True)

    assert result["documents"][0]["errors"] == []
    assert result["documents"][0]["warnings"] == []
    assert result["ok"] is True


def test_form_document_still_runs_header_footer_check(tmp_path, arena_api):
    docx_path = _build_form_docx(tmp_path / "form.docx", header_title="Wrong Title")
    _mock_history(arena_api, [])
    _mock_change_and_item(arena_api)

    result = _verify_one(docx_path, is_form=True)

    assert result["ok"] is False
    assert any("Header Title" in e for e in result["documents"][0]["errors"])


# ---- Sections 2a/2b (unchanged behavior, re-verified against the new fixture) ----


def test_signer_not_a_reviewer_is_blocking(tmp_path, arena_api):
    docx_path = _build_docx(tmp_path / "doc.docx", department_head="Head Person")
    _mock_history(arena_api, ["John Smith", "Head Person"])
    _mock_editions(arena_api)
    _mock_change_and_item(arena_api)

    result = _verify_one(docx_path)

    assert result["ok"] is False
    assert any("Jane Doe" in e for e in result["documents"][0]["errors"])


def test_department_head_not_a_reviewer_is_also_blocking(tmp_path, arena_api):
    docx_path = _build_docx(tmp_path / "doc.docx", department_head="Head Person", signature_fields=[])
    _mock_history(arena_api, [])
    _mock_editions(arena_api)
    _mock_change_and_item(arena_api)

    result = _verify_one(docx_path)

    assert result["ok"] is False
    assert any("Head Person" in e for e in result["documents"][0]["errors"])


def test_preparer_is_not_checked_against_reviewers(tmp_path, arena_api):
    docx_path = _build_docx(
        tmp_path / "doc.docx",
        author="Preparer Person",
        department_head="Jane Doe",
        signature_fields=[],
    )
    _mock_history(arena_api, ["Jane Doe"])
    _mock_editions(arena_api)
    _mock_change_and_item(arena_api)

    result = _verify_one(docx_path)

    assert result["documents"][0]["errors"] == []
    assert not any("Preparer Person" in w for w in result["documents"][0]["warnings"])


def test_eco_date_mismatch_is_a_warning_not_blocking(tmp_path, arena_api):
    docx_path = _build_docx(tmp_path / "doc.docx", department_head="Head Person")
    _mock_history(arena_api, ["Jane Doe", "Head Person"])
    _mock_editions(arena_api, docx_dt="2026-09-02T00:10:00Z", pdf_dt="2026-09-02T00:15:00Z")
    _mock_change_and_item(arena_api)

    result = _verify_one(docx_path)

    doc_result = result["documents"][0]
    assert doc_result["errors"] == []
    assert any("DATE" in w for w in doc_result["warnings"])
    assert result["ok"] is True


def test_reviewer_missing_from_document_is_a_warning_not_blocking(tmp_path, arena_api):
    docx_path = _build_docx(tmp_path / "doc.docx", department_head="Head Person")
    _mock_history(arena_api, ["Jane Doe", "Head Person", "Extra Reviewer"])
    _mock_editions(arena_api)
    _mock_change_and_item(arena_api)

    result = _verify_one(docx_path)

    doc_result = result["documents"][0]
    assert doc_result["errors"] == []
    assert any("extra reviewer" in w.lower() for w in doc_result["warnings"])
    assert result["ok"] is True


def test_document_without_body_tables_is_an_error(tmp_path, arena_api):
    plain = Document()
    section = plain.sections[0]
    htable = section.header.add_table(rows=1, cols=4, width=Emu(0))
    htable.rows[0].cells[1].paragraphs[0].add_run("Title: X")
    htable.rows[0].cells[1].add_paragraph().add_run("Number: FO-00127")
    htable.rows[0].cells[2].paragraphs[0].add_run("Rev 01")
    ftable = section.footer.add_table(rows=2, cols=3, width=Emu(0))
    ftable.rows[0].cells[0].paragraphs[0].add_run("FO-00127 Rev 01")
    ftable.rows[0].cells[1].paragraphs[0].add_run("ECO-00042")
    ftable.rows[0].cells[2].paragraphs[0].add_run(f"Revision Date: {TODAY}")
    plain.add_paragraph("Just some text, no signature/revision tables at all.")
    docx_path = tmp_path / "legacy.docx"
    plain.save(str(docx_path))

    _mock_history(arena_api, [])
    _mock_change_and_item(arena_api)

    result = _verify_one(docx_path, expected_revision="01")

    assert result["ok"] is False
    errors = result["documents"][0]["errors"]
    assert any("signature table" in e for e in errors)
    assert any("revision-history table" in e for e in errors)

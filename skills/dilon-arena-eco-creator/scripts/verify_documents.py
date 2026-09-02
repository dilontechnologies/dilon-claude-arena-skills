#!/usr/bin/env python3
"""Pre-submission document verification -- sections 1, 2a, and 2b.

See docs/requirements/dilon-arena-eco-creator/script-document-verification.md
for the full design record and open questions. Deliberately **does not**
implement section 2c, "Other consistency checks" (still unscoped, not
decided which of its candidates to include -- see that section).

## Source of truth: the compiled .docx, not the markdown source

A Dilon markdown source file isn't always available -- TF (workflow/flow
diagram) documents never have one, and older documents that predate the
markdown-authoring pipeline don't either. What every document is
guaranteed to have is the compiled/uploaded .docx itself, so this script
reads it directly (via python-docx), matching the tables/header/footer
`dilon-document-compiler`/`dilon-document-form-compiler` always produce
(per generate_dilon_doc.py's create_signature_table/create_revision_table
and dilon_docx_common.py's populate_header/populate_footer):

- **Header** (running, every page): a 4-column table -- logo | "Title:
  <title>" + "Number: <doc_number>" (two paragraphs) | "Rev
  <current_revision>" | page N of M.
- **Footer** (running, every page): a 3-column table -- "<doc_number> Rev
  <current_revision>" | "<eco_number>" | "Revision Date: <eco_date>",
  then a merged confidentiality-notice row.
- **Signature table** (body): header row ("Group"/"Preparer"/
  "Signature"), a Preparer row (department/author name), a second header
  row ("Department"/"Name"/"Signature"), a department-head row, then zero
  or more additional-approver rows.
- **Revision history table** (body): a merged "REVISION HISTORY" title
  row, a header row ("REV #"/"DESCRIPTION OF CHANGE"/"ECO #"/"DATE"),
  then one data row per revision. The **last** data row is the current
  revision.

Body tables are located by header row text, not table index, so it's
tolerant of other tables appearing before them in the document. If a
document wasn't produced by either compiler (a legacy document with a
hand-built header/footer/tables in some other shape), these won't be
found -- reported as an error for that document, not guessed around. See
Known gaps in the requirements doc.

## 1. Format compliance: header/footer vs. the actual submission

All of the following are **blocking errors** -- a document whose printed
header/footer doesn't match what's actually being submitted is a
materially wrong record, not a soft warning:
- Header "Title" matches the item's actual Arena name (`get_item`).
- Header "Number" and footer's leading `<doc_number>` both match the
  item's actual Arena number, AND the item number appears somewhere in
  both the header and footer text generally (a looser, redundant check
  independent of the exact-field parse, as a sanity backstop).
- Header "Rev" and footer's `Rev <...>` both match `expected_revision`
  (the `new_revision_number` this item is actually getting on this ECO,
  per SKILL.md step 7 -- passed in, not re-derived here).
- Footer's ECO # matches the actual change's number (`get_change`).
- Footer's Revision Date matches **today's date** (UTC), not an Arena
  upload timestamp -- distinct from 2b below, which checks internal
  consistency between the document and Arena's own upload record; this
  checks the document was actually finalized today, not stale.

## 2a/2b applicability: narrative documents only

Form documents (FO, RE -- compiled via `dilon-document-form-compiler`)
have no signature-approval table and no revision-history table at all --
only a running header/footer. 2a and 2b cannot run against a document
that structurally lacks these tables, so the caller must pass `is_form:
true` for those documents; the script skips both checks entirely for
them (no error, no warning) rather than trying to infer form-vs-narrative
itself (doc-number prefix and table-absence are both unreliable -- see
the requirements doc).

## 2a. Signature table vs. ECO reviewers

Every row below the signature table's second header ("Department"/
"Name"/"Signature") is checked -- both the department-head row and every
additional `signature_fields` row, since the compiled document presents
them identically (this resolves the requirements doc's previously open
question of whether department_head counts: it does, since there's no
structural distinction once rendered). The Preparer row is NOT checked
against ECO reviewers -- preparing a document and formally reviewing it
are different roles.

Each checked name is compared against the ECO's actual reviewers, read
from get_change_history (Arena has no dedicated reviewers endpoint). A
signer who isn't a real ECO reviewer is a **blocking error** -- per the
resolved policy, this should stop progression to "ready to submit." The
reverse direction (an ECO reviewer missing from this document's
signature block) is only a **warning** -- a multi-document ECO may
legitimately have reviewers who aren't required to sign every individual
document. `department` is NOT cross-checked (Arena's user object has no
department field reachable via this API credential).

## 2b. Current-revision date vs. most recent Arena upload

The revision-history table's last row's DATE column is compared against
the later of the docx/pdf file's latest-edition `creationDateTime`
(get_file_editions), using its UTC calendar date. **Warning only, not
blocking** -- timezone handling here is a known, documented
approximation (see Known gaps in the requirements doc): a late-night UTC
upload could land on a different calendar date once converted to Dilon's
actual local timezone, which isn't specified anywhere this script can
read.

This script performs LIVE calls itself (see _arena_import.py for the
required environment variables) and reads local .docx files directly.
Requires python-docx (see env/requirements.txt).

Usage:
    echo '{"change_guid": "CHG123", "documents": [
        {"item_guid": "ITEM1", "docx_path": "C:/.../FO-00127 Rev 02-A.docx",
         "docx_file_guid": "F1", "pdf_file_guid": "F2",
         "expected_revision": "02-A", "is_form": true}
    ]}' | python verify_documents.py

Output: {"documents": [{"item_guid": ..., "errors": [...], "warnings":
         [...]}, ...], "ok": bool}
Exit code 0 if no document has a blocking error, 1 otherwise.
"""
from __future__ import annotations

import json
import re
import sys
from datetime import datetime, timezone
from typing import Any

from docx import Document

from _arena_import import arena

_QUALIFIER_RE = re.compile(r"\s*\([^)]*\)\s*$")

_SIGNATURE_HEADER_1 = ["Group", "Preparer", "Signature"]
_SIGNATURE_HEADER_2 = ["Department", "Name", "Signature"]
_REVISION_HEADER = ["REV #", "DESCRIPTION OF CHANGE", "ECO #", "DATE"]


def _normalize_name(name: str) -> str:
    return re.sub(r"\s+", " ", name).strip().casefold()


def _row_texts(row) -> list[str]:
    return [cell.text.strip() for cell in row.cells]


def _find_table_by_header_row(doc: Document, header_texts: list[str]) -> tuple[Any, int] | None:
    """Return (table, header_row_index) for the first table with a row
    whose leading cell texts exactly match header_texts, or None."""
    for table in doc.tables:
        for row_idx, row in enumerate(table.rows):
            if _row_texts(row)[: len(header_texts)] == header_texts:
                return table, row_idx
    return None


def _extract_signers(doc: Document) -> tuple[list[dict[str, str]], list[str]]:
    """Returns (signers, errors). Signers exclude the Preparer row."""
    found = _find_table_by_header_row(doc, _SIGNATURE_HEADER_1)
    if found is None:
        return [], [
            "Could not find the signature table (expected a 'Group / "
            "Preparer / Signature' header row) -- this document may not "
            "have been compiled via dilon-document-compiler or "
            "dilon-document-form-compiler."
        ]
    table, header1_idx = found
    rows = table.rows
    # Second header ("Department"/"Name"/"Signature") should be the very
    # next row after the Preparer data row.
    header2_idx = header1_idx + 2
    if header2_idx >= len(rows) or _row_texts(rows[header2_idx])[: len(_SIGNATURE_HEADER_2)] != _SIGNATURE_HEADER_2:
        return [], [
            "Signature table found, but its second header row "
            "('Department / Name / Signature') is missing or in an "
            "unexpected position."
        ]

    signers = []
    for row in rows[header2_idx + 1:]:
        texts = _row_texts(row)
        if len(texts) < 2:
            continue
        department, name = texts[0], texts[1]
        if name:
            signers.append({"department": department, "name": name})
    return signers, []


def _extract_current_revision(doc: Document) -> tuple[dict[str, str] | None, list[str]]:
    found = _find_table_by_header_row(doc, _REVISION_HEADER)
    if found is None:
        return None, [
            "Could not find the revision-history table (expected a "
            "'REV # / DESCRIPTION OF CHANGE / ECO # / DATE' header row) "
            "-- this document may not have been compiled via "
            "dilon-document-compiler or dilon-document-form-compiler."
        ]
    table, header_idx = found
    data_rows = table.rows[header_idx + 1:]
    if not data_rows:
        return None, ["Revision-history table found but has no data rows."]
    last = _row_texts(data_rows[-1])
    if len(last) < 4:
        return None, ["Revision-history table's last row is malformed (fewer than 4 cells)."]
    return {"number": last[0], "description": last[1], "eco_number": last[2], "eco_date": last[3]}, []


def _strip_prefix(text: str, prefix: str) -> str:
    return text[len(prefix):].strip() if text.startswith(prefix) else text.strip()


def _extract_header_footer(doc: Document) -> tuple[dict[str, str], dict[str, str], list[str]]:
    errors: list[str] = []
    section = doc.sections[0]

    header_fields: dict[str, str] = {}
    header_tables = section.header.tables
    if not header_tables:
        errors.append(
            "No header table found -- this document may not have been "
            "compiled via dilon-document-compiler or "
            "dilon-document-form-compiler."
        )
    else:
        row = header_tables[0].rows[0]
        title_cell_lines = [p.text.strip() for p in row.cells[1].paragraphs if p.text.strip()]
        title = next((_strip_prefix(t, "Title:") for t in title_cell_lines if t.startswith("Title:")), None)
        number = next((_strip_prefix(t, "Number:") for t in title_cell_lines if t.startswith("Number:")), None)
        header_fields = {
            "title": title,
            "number": number,
            "revision": _strip_prefix(row.cells[2].text.strip(), "Rev"),
            "full_text": " ".join(c.text for c in row.cells),
        }

    footer_fields: dict[str, str] = {}
    footer_tables = section.footer.tables
    if not footer_tables:
        errors.append(
            "No footer table found -- this document may not have been "
            "compiled via dilon-document-compiler or "
            "dilon-document-form-compiler."
        )
    else:
        row = footer_tables[0].rows[0]
        footer_fields = {
            "id_line": row.cells[0].text.strip(),
            "eco_number": row.cells[1].text.strip(),
            "revision_date": _strip_prefix(row.cells[2].text.strip(), "Revision Date:"),
            "full_text": " ".join(c.text for c in row.cells),
        }

    return header_fields, footer_fields, errors


def _check_header_footer(
    header_fields: dict[str, str],
    footer_fields: dict[str, str],
    expected_item_number: str,
    expected_revision: str,
    expected_eco_number: str,
    expected_title: str,
    today_iso: str,
) -> list[str]:
    errors: list[str] = []

    if header_fields.get("title") != expected_title:
        errors.append(
            f"Header Title {header_fields.get('title')!r} does not match "
            f"the item's actual name {expected_title!r}."
        )
    if header_fields.get("number") != expected_item_number:
        errors.append(
            f"Header Number {header_fields.get('number')!r} does not match "
            f"the item's actual number {expected_item_number!r}."
        )
    if header_fields.get("revision") != expected_revision:
        errors.append(
            f"Header Rev {header_fields.get('revision')!r} does not match "
            f"the expected revision {expected_revision!r}."
        )
    if expected_item_number not in (header_fields.get("full_text") or ""):
        errors.append(f"Item number {expected_item_number!r} does not appear anywhere in the header.")

    expected_id_line = f"{expected_item_number} Rev {expected_revision}"
    if footer_fields.get("id_line") != expected_id_line:
        errors.append(
            f"Footer ID line {footer_fields.get('id_line')!r} does not "
            f"match expected {expected_id_line!r}."
        )
    if footer_fields.get("eco_number") != expected_eco_number:
        errors.append(
            f"Footer ECO # {footer_fields.get('eco_number')!r} does not "
            f"match the actual ECO number {expected_eco_number!r}."
        )
    if footer_fields.get("revision_date") != today_iso:
        errors.append(
            f"Footer Revision Date {footer_fields.get('revision_date')!r} "
            f"does not match today's date {today_iso!r}."
        )
    if expected_item_number not in (footer_fields.get("full_text") or ""):
        errors.append(f"Item number {expected_item_number!r} does not appear anywhere in the footer.")

    return errors


def _eco_reviewer_names(change_guid: str) -> tuple[set[str], list[str]]:
    """Returns (normalized reviewer names, errors)."""
    history = arena.get_change_history(change_guid)
    if not isinstance(history, dict) or history.get("error"):
        return set(), [f"get_change_history failed: {history!r}"]

    names: set[str] = set()
    for entry in history.get("results", []) or []:
        if entry.get("action") != "Add Additional Reviewer":
            continue
        raw = entry.get("newValue") or ""
        stripped = _QUALIFIER_RE.sub("", raw)
        if stripped:
            names.add(_normalize_name(stripped))
    return names, []


def _check_signatures(signers: list[dict[str, str]], eco_reviewers: set[str]) -> tuple[list[str], list[str]]:
    errors: list[str] = []
    warnings: list[str] = []
    doc_signers = set()
    for entry in signers:
        norm = _normalize_name(entry["name"])
        doc_signers.add(norm)
        if norm not in eco_reviewers:
            errors.append(f"Signer {entry['name']!r} is not a recorded reviewer on this ECO.")

    for reviewer in eco_reviewers - doc_signers:
        warnings.append(
            f"ECO reviewer (normalized) {reviewer!r} does not appear in this "
            "document's signature table."
        )
    return errors, warnings


def _latest_edition_datetime(file_guid: str) -> datetime | None:
    editions = arena.get_file_editions(file_guid)
    if not isinstance(editions, dict) or editions.get("error"):
        return None
    best: datetime | None = None
    for edition in editions.get("results", []) or []:
        raw = edition.get("creationDateTime")
        if not raw:
            continue
        dt = datetime.fromisoformat(raw.replace("Z", "+00:00"))
        if best is None or dt > best:
            best = dt
    return best


def _check_eco_date(
    current_revision: dict[str, str] | None, docx_file_guid: str | None, pdf_file_guid: str | None
) -> list[str]:
    warnings: list[str] = []
    eco_date = current_revision.get("eco_date") if current_revision else None
    if not eco_date:
        warnings.append("No eco_date available from the revision-history table's last row.")
        return warnings

    candidates = [
        dt
        for dt in (
            _latest_edition_datetime(g) for g in (docx_file_guid, pdf_file_guid) if g
        )
        if dt is not None
    ]
    if not candidates:
        warnings.append("Could not read any file edition timestamp to compare against eco_date.")
        return warnings

    latest = max(candidates).astimezone(timezone.utc).date().isoformat()
    if latest != eco_date:
        warnings.append(
            f"Revision table's DATE {eco_date!r} does not match the latest Arena "
            f"upload's UTC date {latest!r} -- timezone-naive comparison, see "
            "script docstring."
        )
    return warnings


def _fetch_item_number_and_name(item_guid: str) -> tuple[str | None, str | None, list[str]]:
    item = arena.get_item(item_guid)
    if not isinstance(item, dict) or item.get("error"):
        return None, None, [f"get_item({item_guid!r}) failed: {item!r}"]
    return item.get("number"), item.get("name"), []


def verify(change_guid: str, documents: list[dict[str, Any]]) -> dict[str, Any]:
    eco_reviewers, history_errors = _eco_reviewer_names(change_guid)

    change = arena.get_change(change_guid)
    eco_number = change.get("number") if isinstance(change, dict) else None
    change_errors = [] if eco_number else [f"get_change({change_guid!r}) failed or has no number: {change!r}"]

    today_iso = datetime.now(timezone.utc).date().isoformat()

    results = []
    any_error = bool(history_errors) or bool(change_errors)

    for doc_input in documents:
        item_guid = doc_input["item_guid"]
        docx_path = doc_input["docx_path"]
        expected_revision = doc_input["expected_revision"]
        is_form = doc_input["is_form"]
        errors = list(history_errors) + list(change_errors)
        warnings: list[str] = []

        doc = Document(docx_path)

        item_number, item_name, item_errors = _fetch_item_number_and_name(item_guid)
        errors.extend(item_errors)

        header_fields, footer_fields, hf_errors = _extract_header_footer(doc)
        errors.extend(hf_errors)
        if not hf_errors and not item_errors and eco_number:
            errors.extend(
                _check_header_footer(
                    header_fields,
                    footer_fields,
                    item_number,
                    expected_revision,
                    eco_number,
                    item_name,
                    today_iso,
                )
            )

        if not is_form:
            signers, sig_extract_errors = _extract_signers(doc)
            errors.extend(sig_extract_errors)
            if not sig_extract_errors:
                sig_errors, sig_warnings = _check_signatures(signers, eco_reviewers)
                errors.extend(sig_errors)
                warnings.extend(sig_warnings)

            current_revision, rev_extract_errors = _extract_current_revision(doc)
            errors.extend(rev_extract_errors)
            if not rev_extract_errors:
                warnings.extend(
                    _check_eco_date(current_revision, doc_input.get("docx_file_guid"), doc_input.get("pdf_file_guid"))
                )

        if errors:
            any_error = True
        results.append({"item_guid": item_guid, "errors": errors, "warnings": warnings})

    return {"documents": results, "ok": not any_error}


def main() -> int:
    try:
        payload = json.load(sys.stdin)
        result = verify(payload["change_guid"], payload["documents"])
    except (json.JSONDecodeError, KeyError, TypeError) as exc:
        print(json.dumps({"error": True, "message": f"Bad input: {exc}"}))
        return 1
    print(json.dumps(result))
    return 0 if result.get("ok") else 1


if __name__ == "__main__":
    raise SystemExit(main())
